"""Saved evidence investigations; no model or risk policy is changed."""
import io, json, math
from datetime import datetime, timezone
from html import escape
from statistics import mean, pstdev
from fastapi import Depends, HTTPException
from fastapi.responses import Response


def baseline(reference, observed, features):
    output = []
    for feature in features:
        values = [r['feature_evidence'][feature] for r in reference if feature in r.get('feature_evidence', {})]
        if len(values) < 20:
            output.append({'feature': feature, 'status': 'insufficient_history', 'samples': len(values)})
            continue
        avg, sd = mean(values), pstdev(values)
        value = observed['feature_evidence'][feature]
        output.append({'feature': feature, 'status': 'available', 'samples': len(values),
                       'observed': value, 'reference_mean': avg, 'reference_std': sd,
                       'z_score': (value-avg)/sd if sd else None,
                       'constant_reference_deviation': bool(sd == 0 and value != avg)})
    return output


def spreadsheet(rows):
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = 'Incidents'
    keys = ['id','scan_id','src_ip','dst_ip','label','risk','status','occurrences','created','note']
    ws.append(keys)
    for row in rows:
        # Force string cells to text so analyst notes cannot become Excel formulas.
        ws.append([row.get(k) for k in keys])
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, str): cell.data_type = 's'
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(55,max(16,len(str(col[0].value))+3))
    f=io.BytesIO(); wb.save(f); return f.getvalue()


def case_pdf(case):
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.pagesizes import A4
    f=io.BytesIO(); styles=getSampleStyleSheet(); story=[]
    def p(text, style='BodyText'): story.append(Paragraph(escape(str(text)),styles[style]));story.append(Spacer(1,8))
    incident=case['incident']
    p('TraceVigil investigation case '+str(incident['id']),'Title')
    p('Generated UTC '+datetime.now(timezone.utc).isoformat())
    for k in ['scan_id','src_ip','dst_ip','label','risk','status','occurrences','note']:p(k+': '+str(incident.get(k)))
    p('Interpretation','Heading2');p('Model predictions are investigation leads, not confirmed compromise. Risk is a policy priority, not a probability. IP addresses do not establish a person’s identity.')
    p('Observed flow evidence','Heading2')
    for r in case['flows'][:10]:
        p('Row '+str(r['row_number'])+' | '+str(r.get('timestamp'))+' | '+str(r.get('consensus'))+' | priority '+str(r.get('risk_score')))
        p('Review reasons: '+', '.join(r.get('review_reasons',[])))
    p('Evidence coverage','Heading2');p(case['coverage'])
    p('Historical reference','Heading2');p(case['baseline_scope'])
    for r in case['baseline']:
        p(r['feature']+': '+ ('insufficient history ('+str(r['samples'])+' records)' if r['status']!='available' else 'observed '+str(r['observed'])+'; historical mean '+str(round(r['reference_mean'],4))+'; z score '+str(r['z_score'])))
    if case.get('explanation'):
        explanation=case['explanation'];p('Local XGBoost TreeSHAP explanation','Heading2')
        p(explanation['scope']);p('Class: '+explanation['class']+'; units: '+explanation['units'])
        p('Base value '+str(explanation['base_value'])+'; raw margin '+str(explanation['margin'])+'; additive check '+str(explanation['sum_check']))
        for row in explanation['features']:p(row['feature']+': '+str(round(row['contribution'],5)))
    p('Analyst next steps','Heading2')
    for step in case['next_steps']:p(step)
    SimpleDocTemplate(f,pagesize=A4,rightMargin=40,leftMargin=40).build(story)
    return f.getvalue()


def install(m):
    app=m.app
    def get_case(incident,u):
        with m.connection() as c:
            row=c.execute('SELECT * FROM incidents WHERE id=?',(incident,)).fetchone()
            if row is None:raise HTTPException(404,'Incident not found')
            item=dict(row);m.authorize(item['scan_id'],u)
            # Bound expensive JSON processing and explicitly disclose truncation.
            records=c.execute('SELECT payload FROM flows WHERE scan_id=? ORDER BY row_number LIMIT 50001',(item['scan_id'],)).fetchall()
            saved_anchor=c.execute('SELECT payload FROM flows WHERE scan_id=? AND row_number=?',(item['scan_id'],item['row_number'])).fetchone()
        truncated=len(records)>50000
        rows=[json.loads(r['payload']) for r in records[:50000]]
        related=[r for r in rows if r.get('status')=='scored' and
                 (r['src_ip']['value'] or 'Unavailable')==item['src_ip'] and
                 (r['dst_ip']['value'] or 'Unavailable')==item['dst_ip'] and r.get('consensus')==item['label'] and r.get('needs_review')]
        anchor=json.loads(saved_anchor['payload']) if saved_anchor else None
        if anchor and not any(r['row_number']==anchor['row_number'] for r in related):related.insert(0,anchor)
        reference=[]
        if anchor and anchor['src_ip'].get('status')=='valid':
            import pandas as pd
            stamp=pd.to_datetime(anchor.get('timestamp'),errors='coerce',dayfirst=True,format='mixed')
            if not pd.isna(stamp):
                for r in rows:
                    if r.get('status')!='scored' or r['row_number']>=anchor['row_number'] or r['src_ip']['value']!=anchor['src_ip']['value']:continue
                    earlier=pd.to_datetime(r.get('timestamp'),errors='coerce',dayfirst=True,format='mixed')
                    try:
                        if not pd.isna(earlier) and earlier<stamp:reference.append(r)
                    except TypeError:pass
        return {'incident':item,'flows':related[:200], 'anchor':anchor,
                'baseline':baseline(reference,anchor,m.engine.features) if anchor else [],
                'baseline_scope':'Same scan and source IP, earlier row AND earlier capture timestamp than the first incident row; minimum 20 records per feature. Descriptive history may contain attacks, is not a verified benign baseline, and never changes risk. Missing/unparseable timestamps yield insufficient history.',
                'coverage':f'Inspected first {len(rows)} saved rows plus direct first-incident-row lookup; window capped at 50000; truncated={truncated}. Showing at most {min(200,len(related))} matching review flows from this window and anchor. PDF shows at most 10 flows.',
                'next_steps':['Validate the CSV source, capture timing and endpoint ownership.','Compare the three model predictions and review uncertainty.','Check authorised activity and corroborating network or endpoint logs.','Record findings before considering containment; this tool does not block traffic.']}

    @app.get('/api/overview')
    def overview(u=Depends(m.user)):
        clause='' if u['role']=='Admin' else ' WHERE owner=?';args=() if not clause else (u['username'],)
        with m.connection() as c:
            scans=c.execute('SELECT status,summary FROM scans'+clause,args).fetchall()
            incidents=c.execute('SELECT status,risk,occurrences FROM incidents'+clause,args).fetchall()
        totals={'scans':len(scans),'rows_read':0,'scored':0,'open_incidents':sum(x['status'] in ('Open','Investigating') for x in incidents),'critical_groups':sum(x['risk']>=76 for x in incidents),'review_flows':0}
        labels={};risks={}
        for s in scans:
            d=json.loads(s['summary']) if s['summary'] else {}
            totals['rows_read']+=d.get('rows_read',0);totals['scored']+=d.get('scored',0);totals['review_flows']+=d.get('alerts',0)
            for k,v in d.get('attack_types',{}).items():labels[k]=labels.get(k,0)+v
            for k,v in d.get('risk_distribution',{}).items():risks[k]=risks.get(k,0)+v
        return {'totals':totals,'traffic_types':labels,'risk_distribution':risks,'scope':'Saved scan summaries visible to your role. Repeated scans count again; running scans may have no saved final summary. Critical groups include all statuses; none are automatically confirmed attacks.'}

    @app.get('/api/incidents-export.xlsx')
    def xlsx(u=Depends(m.user)):
        clause='' if u['role']=='Admin' else ' WHERE owner=?';args=() if not clause else (u['username'],)
        with m.connection() as c:rows=[dict(r) for r in c.execute('SELECT * FROM incidents'+clause+' ORDER BY id DESC LIMIT 10000',args)]
        return Response(spreadsheet(rows),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':'attachment; filename="tracevigil-incidents-latest-10000.xlsx"'})

    @app.get('/api/incidents/{incident}/case')
    def case(incident:int,u=Depends(m.user)):return get_case(incident,u)

    @app.get('/api/incidents/{incident}/case.pdf')
    def pdf(incident:int,u=Depends(m.user)):
        data=get_case(incident,u)
        if data['anchor']:data['explanation']=local_explanation(data['anchor'],m)
        return Response(case_pdf(data),media_type='application/pdf',headers={'Content-Disposition':f'attachment; filename="tracevigil-case-{incident}.pdf"'})

    @app.get('/api/incidents/{incident}/explanation')
    def explanation(incident:int,u=Depends(m.user)):
        case=get_case(incident,u);anchor=case['anchor']
        if not anchor:raise HTTPException(422,'Saved first incident row is unavailable')
        return local_explanation(anchor,m)


def local_explanation(anchor,m):
    import numpy as np
    import xgboost as xgb
    model=m.engine.models['xgboost'];features=m.engine.features
    matrix=xgb.DMatrix(np.array([[anchor['feature_evidence'][f] for f in features]]),feature_names=features)
    booster=model.get_booster()
    contributions=booster.predict(matrix,pred_contribs=True,strict_shape=True)[0]
    margins=booster.predict(matrix,output_margin=True,strict_shape=True)[0]
    idx=int(np.argmax(margins));values=contributions[idx]
    label=str(m.engine.bundle['encoder'].inverse_transform([idx])[0])
    return {'model':'xgboost','class':label,'units':'raw class margin (before softmax), NOT probability or risk points','base_value':float(values[-1]),'margin':float(margins[idx]),'sum_check':bool(np.isclose(values.sum(),margins[idx],atol=1e-4)), 'features':sorted([{'feature':f,'observed':anchor['feature_evidence'][f],'contribution':float(v)} for f,v in zip(features,values[:-1])],key=lambda r:abs(r['contribution']),reverse=True),'scope':'Native XGBoost TreeSHAP for the first incident row and its XGBoost-selected class. Explains XGBoost only, not majority vote, policy risk or causal responsibility.'}
